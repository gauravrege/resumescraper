"""
Excel Exporter Module for Resume Scraper
Creates sleek, recruiter-friendly multi-tab ATS workbooks with sheet navigation,
candidate status dropdowns, and decluttered layouts.
"""

import os
from datetime import datetime
from typing import List, Dict, Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from config import ATS_STATUSES


# Style definitions
COLOR_PRIMARY_DARK = "0F172A"   # Slate 900
COLOR_PRIMARY = "1E293B"        # Slate 800
COLOR_HEADER_BG = "1E293B"      # Dark slate header
COLOR_ZEBRA = "F8FAFC"          # Slate 50
COLOR_BORDER = "E2E8F0"         # Slate 200
COLOR_NAV_BG = "F1F5F9"         # Slate 100
COLOR_CARD_BG = "F8FAFC"        # Card background
COLOR_ACCENT = "0284C7"         # Sky 600

FONT_NAME = "Segoe UI"
FONT_HEADER = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
FONT_DATA = Font(name=FONT_NAME, size=10)
FONT_BOLD = Font(name=FONT_NAME, size=10, bold=True)
FONT_LINK = Font(name=FONT_NAME, size=10, color="0284C7", underline="single")
FONT_NAV = Font(name=FONT_NAME, size=10, bold=True, color="0F172A")

FILL_HEADER = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
FILL_ZEBRA = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid")
FILL_NAV = PatternFill(start_color=COLOR_NAV_BG, end_color=COLOR_NAV_BG, fill_type="solid")

BORDER_THIN = Border(
    left=Side(style='thin', color=COLOR_BORDER),
    right=Side(style='thin', color=COLOR_BORDER),
    top=Side(style='thin', color=COLOR_BORDER),
    bottom=Side(style='thin', color=COLOR_BORDER)
)


def apply_table_styling(ws, header_row: int, is_master: bool = False):
    """Styles headers, data rows, zebra striping, and auto-adjusts column widths."""
    ws.views.sheetView[0].showGridLines = True
    max_col = ws.max_column
    max_row = ws.max_row

    # Auto-filter and header row height
    max_col_letter = get_column_letter(max_col)
    ws.auto_filter.ref = f"A{header_row}:{max_col_letter}{max_row}"
    ws.row_dimensions[header_row].height = 26

    # Style header row
    for col_idx in range(1, max_col + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER_THIN

    # Style data rows
    for row_idx in range(header_row + 1, max_row + 1):
        ws.row_dimensions[row_idx].height = 22
        is_zebra = (row_idx % 2 == 1)
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = BORDER_THIN
            if not cell.font.underline:
                cell.font = FONT_DATA
            cell.alignment = Alignment(vertical="center")
            if is_zebra and not cell.fill.start_color.rgb:
                cell.fill = FILL_ZEBRA

    # Auto-fit column widths with sensible min/max limits
    for col in ws.iter_cols(min_row=header_row, max_row=max_row, min_col=1, max_col=max_col):
        col_letter = get_column_letter(col[0].column)
        header_val = str(ws.cell(row=header_row, column=col[0].column).value or "")
        max_len = len(header_val)
        for cell in col[1:]:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        # Bounded between 13 and 38 for readability
        adjusted_width = min(max(max_len + 3, 13), 38)
        # Dedicated column width adjustments
        if "Resume" in header_val:
            adjusted_width = 14
        elif "Status" in header_val:
            adjusted_width = 20
        elif "Top Skills" in header_val:
            adjusted_width = 32
        elif "Notes" in header_val:
            adjusted_width = 24
        ws.column_dimensions[col_letter].width = adjusted_width


def add_status_validation(ws, start_row: int, status_col_idx: int):
    """Adds ATS candidate status dropdown validation."""
    if start_row > ws.max_row:
        return
    col_letter = get_column_letter(status_col_idx)
    formula = f'"{",".join(ATS_STATUSES)}"'
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{start_row}:{col_letter}{ws.max_row}")


def build_dashboard(wb, records: List[Dict[str, Any]], sheet_mapping: Dict[str, str]):
    """Creates a modern KPI Summary Dashboard."""
    ws = wb.create_sheet(title="Summary Dashboard", index=0)
    ws.views.sheetView[0].showGridLines = True

    # Title & Subtitle
    ws["B2"] = "TALENT ACQUISITION & RESUME DASHBOARD"
    ws["B2"].font = Font(name=FONT_NAME, size=16, bold=True, color=COLOR_PRIMARY_DARK)
    ws["B3"] = f"Report Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}  |  Firm Talent Pipeline"
    ws["B3"].font = Font(name=FONT_NAME, size=10, italic=True, color="64748B")

    # Metrics
    total_count = len(records)
    processed_count = sum(1 for r in records if r.get("Status") == "Processed")
    domain_count = len(sheet_mapping)

    # KPI Card 1: Total Candidates
    ws["B5"] = "TOTAL CANDIDATES"
    ws["B5"].font = Font(name=FONT_NAME, size=9, bold=True, color="64748B")
    ws["B6"] = total_count
    ws["B6"].font = Font(name=FONT_NAME, size=20, bold=True, color=COLOR_PRIMARY_DARK)

    # KPI Card 2: Segregated Domains
    ws["D5"] = "TALENT DOMAINS"
    ws["D5"].font = Font(name=FONT_NAME, size=9, bold=True, color="64748B")
    ws["D6"] = domain_count
    ws["D6"].font = Font(name=FONT_NAME, size=20, bold=True, color=COLOR_PRIMARY_DARK)

    # KPI Card 3: Processing Rate
    ws["F5"] = "SUCCESS RATE"
    ws["F5"].font = Font(name=FONT_NAME, size=9, bold=True, color="64748B")
    rate = f"{(processed_count / total_count * 100):.0f}%" if total_count else "0%"
    ws["F6"] = rate
    ws["F6"].font = Font(name=FONT_NAME, size=20, bold=True, color="16A34A")

    # Style KPI boxes
    card_fill = PatternFill(start_color=COLOR_CARD_BG, end_color=COLOR_CARD_BG, fill_type="solid")
    for start_col in ["B", "D", "F"]:
        col_end = chr(ord(start_col) + 1)
        ws.merge_cells(f"{start_col}5:{col_end}5")
        ws.merge_cells(f"{start_col}6:{col_end}6")
        for r in [5, 6]:
            for c in [start_col, col_end]:
                cell = ws[f"{c}{r}"]
                cell.fill = card_fill
                cell.border = BORDER_THIN
                cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[5].height = 18
    ws.row_dimensions[6].height = 32

    # Master Overview Quick Link
    ws["B8"] = "Quick Navigation"
    ws["B8"].font = Font(name=FONT_NAME, size=12, bold=True, color=COLOR_PRIMARY_DARK)

    ws["B9"] = "📋 Master Sheet (All Candidates)"
    ws["B9"].font = FONT_LINK
    ws["B9"].hyperlink = "#'All Candidates'!A1"

    # Category Breakdown Table
    ws["B11"] = "Domain Category Breakdown"
    ws["B11"].font = Font(name=FONT_NAME, size=12, bold=True, color=COLOR_PRIMARY_DARK)

    table_headers = ["Domain / Department", "Candidates", "% of Pool", "Quick Action"]
    cols = ["B", "C", "D", "E"]
    for c, h in zip(cols, table_headers):
        cell = ws[f"{c}12"]
        cell.value = h
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.border = BORDER_THIN
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[12].height = 24

    # Calculate category counts
    cat_counts: Dict[str, int] = {}
    for r in records:
        cat = r.get("Category", "General_and_Others")
        cat_counts[cat] = cat_counts.get(cat, 0) + 1

    sorted_cats = sorted(cat_counts.items(), key=lambda x: x[1], reverse=True)
    curr_row = 13
    for cat, count in sorted_cats:
        sheet_title = sheet_mapping.get(cat, cat.replace("_", " ")[:31])
        pct = f"{(count / total_count * 100):.1f}%" if total_count else "0%"

        ws[f"B{curr_row}"] = cat.replace("_", " ")
        ws[f"C{curr_row}"] = count
        ws[f"D{curr_row}"] = pct
        ws[f"E{curr_row}"] = f"View {sheet_title} ➔"

        # Hyperlink to category sheet
        ws[f"B{curr_row}"].hyperlink = f"#'{sheet_title}'!A1"
        ws[f"B{curr_row}"].font = FONT_BOLD
        ws[f"C{curr_row}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"D{curr_row}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"E{curr_row}"].hyperlink = f"#'{sheet_title}'!A1"
        ws[f"E{curr_row}"].font = FONT_LINK
        ws[f"E{curr_row}"].alignment = Alignment(horizontal="center", vertical="center")

        is_zebra = (curr_row % 2 == 1)
        for col_letter in cols:
            cell = ws[f"{col_letter}{curr_row}"]
            cell.border = BORDER_THIN
            if is_zebra:
                cell.fill = FILL_ZEBRA

        ws.row_dimensions[curr_row].height = 22
        curr_row += 1

    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 24
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 16


def export_to_excel(records: List[Dict[str, Any]], output_path: str) -> str:
    """
    Exports parsed resume records into a decluttered, multi-tab ATS workbook.
    """
    if not records:
        return ""

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # 1. Prepare Master Sheet ("All Candidates")
    master_ws = wb.create_sheet(title="All Candidates")
    master_ws.freeze_panes = "A2"

    master_columns = [
        "Candidate Name",
        "Domain Category",
        "Designation",
        "Current Company",
        "Location",
        "Experience",
        "Education",
        "Top Skills",
        "Email",
        "Phone",
        "Candidate Status",
        "Recruiter Notes",
        "Resume"
    ]

    # Write Master Headers
    for c_idx, col_name in enumerate(master_columns, start=1):
        master_ws.cell(row=1, column=c_idx, value=col_name)

    # Write Master Rows
    for r_idx, rec in enumerate(records, start=2):
        clean_cat = rec.get("Category", "General_and_Others").replace("_", " ")
        master_ws.cell(row=r_idx, column=1, value=rec.get("Candidate Name"))
        master_ws.cell(row=r_idx, column=2, value=clean_cat)
        master_ws.cell(row=r_idx, column=3, value=rec.get("Designation"))
        master_ws.cell(row=r_idx, column=4, value=rec.get("Current Company"))
        master_ws.cell(row=r_idx, column=5, value=rec.get("Location"))
        master_ws.cell(row=r_idx, column=6, value=rec.get("Experience"))
        master_ws.cell(row=r_idx, column=7, value=rec.get("Education"))
        master_ws.cell(row=r_idx, column=8, value=rec.get("Top Skills"))
        master_ws.cell(row=r_idx, column=9, value=rec.get("Email"))
        master_ws.cell(row=r_idx, column=10, value=rec.get("Phone"))
        master_ws.cell(row=r_idx, column=11, value=rec.get("Candidate Status", "New"))
        master_ws.cell(row=r_idx, column=12, value=rec.get("Recruiter Notes", ""))

        # Clean compact resume link
        cv_cell = master_ws.cell(row=r_idx, column=13, value="📄 Open CV")
        file_p = rec.get("File Path", "")
        if file_p and os.path.exists(file_p):
            cv_cell.hyperlink = f"file:///{os.path.abspath(file_p).replace(chr(92), '/')}"
            cv_cell.font = FONT_LINK
            cv_cell.alignment = Alignment(horizontal="center", vertical="center")

    apply_table_styling(master_ws, header_row=1, is_master=True)
    add_status_validation(master_ws, start_row=2, status_col_idx=11)

    # 2. Segregate into Category Tabs
    category_columns = [
        "Candidate Name",
        "Designation",
        "Current Company",
        "Location",
        "Experience",
        "Education",
        "Top Skills",
        "Email",
        "Phone",
        "Candidate Status",
        "Recruiter Notes",
        "Resume"
    ]

    # Map categories to sheet titles (max 31 chars)
    unique_categories = sorted(list(set(r.get("Category", "General_and_Others") for r in records)))
    sheet_mapping: Dict[str, str] = {}
    for cat in unique_categories:
        clean_title = cat.replace("_and_", " & ").replace("_", " ")[:31]
        sheet_mapping[cat] = clean_title

    for cat in unique_categories:
        sheet_title = sheet_mapping[cat]
        cat_ws = wb.create_sheet(title=sheet_title)
        cat_ws.freeze_panes = "A4"

        # Navigation Bar in Row 1 & 2
        cat_ws["A1"] = "⬅ Back to Dashboard"
        cat_ws["A1"].font = FONT_NAV
        cat_ws["A1"].hyperlink = "#'Summary Dashboard'!A1"
        cat_ws["A1"].fill = FILL_NAV
        cat_ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        cat_ws.merge_cells("A1:B1")
        cat_ws.row_dimensions[1].height = 22

        cat_candidates = [r for r in records if r.get("Category") == cat]
        cat_ws["A2"] = f"Domain: {cat.replace('_', ' ')}  |  {len(cat_candidates)} Candidate(s) Found"
        cat_ws["A2"].font = Font(name=FONT_NAME, size=10, italic=True, color="64748B")
        cat_ws.row_dimensions[2].height = 18

        # Headers in Row 3
        for c_idx, col_name in enumerate(category_columns, start=1):
            cat_ws.cell(row=3, column=c_idx, value=col_name)

        # Data rows in Row 4+
        for r_idx, rec in enumerate(cat_candidates, start=4):
            cat_ws.cell(row=r_idx, column=1, value=rec.get("Candidate Name"))
            cat_ws.cell(row=r_idx, column=2, value=rec.get("Designation"))
            cat_ws.cell(row=r_idx, column=3, value=rec.get("Current Company"))
            cat_ws.cell(row=r_idx, column=4, value=rec.get("Location"))
            cat_ws.cell(row=r_idx, column=5, value=rec.get("Experience"))
            cat_ws.cell(row=r_idx, column=6, value=rec.get("Education"))
            cat_ws.cell(row=r_idx, column=7, value=rec.get("Top Skills"))
            cat_ws.cell(row=r_idx, column=8, value=rec.get("Email"))
            cat_ws.cell(row=r_idx, column=9, value=rec.get("Phone"))
            cat_ws.cell(row=r_idx, column=10, value=rec.get("Candidate Status", "New"))
            cat_ws.cell(row=r_idx, column=11, value=rec.get("Recruiter Notes", ""))

            cv_cell = cat_ws.cell(row=r_idx, column=12, value="📄 Open CV")
            file_p = rec.get("File Path", "")
            if file_p and os.path.exists(file_p):
                cv_cell.hyperlink = f"file:///{os.path.abspath(file_p).replace(chr(92), '/')}"
                cv_cell.font = FONT_LINK
                cv_cell.alignment = Alignment(horizontal="center", vertical="center")

        apply_table_styling(cat_ws, header_row=3, is_master=False)
        add_status_validation(cat_ws, start_row=4, status_col_idx=10)

    # 3. Create Dashboard at Sheet Index 0
    build_dashboard(wb, records, sheet_mapping)

    wb.save(output_path)
    return output_path
